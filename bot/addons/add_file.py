import os
import logging
from docx import Document
from openpyxl import workbook
from bot.database import db

import datetime
import openpyxl


# Открывает документ который лежит в папке donwloads
def open_file_docx(name_file: str, company_pincode: list):
    basedir = os.path.join("downloads", name_file)
    doc = Document(basedir)
    dict_quest = dict()
    # Читает документ
    try:
        for paragraph in doc.paragraphs:
            if paragraph.text:
                # Если в документе строка с ** то записывает
                # эту строку как пинкод
                if "**" in paragraph.text:
                    pincode = paragraph.text.replace("**", "")
                # Если в документе строка с * то записывает
                # эту строку как вакансию
                if "*" in paragraph.text:
                    vac = paragraph.text.replace("*", "")
                    num = 1
                    vacan = dict()
                    dict_quest["slots"] = {vac: vacan}
                # Строки без метки становятся вопросами
                else:
                    questions = (paragraph.text.replace("\n", "")).strip()
                    vacan[str(num)] = questions
                    num += 1
                # Формирует словарь с данными
                dict_quest["date"] = datetime.datetime.now()
                dict_quest["pincode"] = (str(pincode)).strip()
                dict_quest["slots"][vac] = vacan
        text_msg = add_quest_vacan_slot(dict_quest, company_pincode)
    except UnboundLocalError as err:
        text_msg = 'Вы забыли поставить метку около вакансии "*"'
        logging.error(f"Ошибка {err} | from add_file")
    return text_msg


def excel_read_blank(path: str, company_pincode: list[str]):
    path_chain = path.split("_")
    dict_quest = {
        "date": datetime.datetime.now(),
        "pincode": path_chain[-2],
        "slots": {},
    }
    base_path = os.path.join("downloads", path)
    workbook = openpyxl.load_workbook(base_path)
    dict_quest["slots"][path_chain[-1][:-5]] = {
        str(indx + 1): str(row[0].value or "").replace(";", "\\\\n")
        for indx, row in enumerate(workbook.active)
        if row[0].value
    }
    text_msg = add_quest_vacan_slot(dict_quest, company_pincode)
    return text_msg


def add_quest_vacan_slot(info: dict, company_pincode: list):
    format_add_db_dict = create_dict_db(company_pincode, info)
    # Добавляет данные в базу или нет если есть совпадение по пинкоду
    answer_db = db.get_or_create_job(db.db, format_add_db_dict)
    if not answer_db:
        return "Совпадение пинкодов данные не добавлены"
    return "Данные успешно добавлены"


# Cоздает итоговый словарь для оправки в базу
def create_dict_db(company_pincode: list, blank: dict):
    db_dict = {
        "company": company_pincode[0],
        "vacancy": blank,
        "pincode": company_pincode[1], 
    }
    return db_dict
